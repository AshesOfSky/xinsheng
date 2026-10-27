import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook, load_workbook
from private_file_worker import FileTaskError, is_review_request, run_file_task


class ReviewValidationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)

    def run_task(self, source, output, request):
        def runner(folder, request, filename):
            shutil.copy2(output, folder / 'output' / ('result' + source.suffix))
        with patch('private_file_worker.locate_attachment', return_value=source), \
             patch('private_file_worker.check_text'):
            return run_file_task(None, 'wxid_friend', {}, request, None, self.root, runner=runner)

    def xlsx(self, filename, report=None):
        book = Workbook()
        book.active['A1'] = '毛利表'
        if report is not None:
            sheet = book.create_sheet('核查结果')
            sheet['A1'] = report
        path = self.root / filename
        book.save(path)
        return path

    def docx(self, filename, paragraphs):
        document = Document()
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
        path = self.root / filename
        document.save(path)
        return path

    def test_xlsx_review_requires_nonempty_report(self):
        source = self.xlsx('source.xlsx')
        for value in (None, ''):
            with self.subTest(report=value):
                output = self.xlsx('output.xlsx', report=value)
                with self.assertRaisesRegex(FileTaskError, '核查结果'):
                    self.run_task(source, output, '检查一下毛利表')
        output = self.xlsx('output.xlsx', report='检查范围：收入与成本公式')
        self.assertTrue(self.run_task(source, output, '检查一下毛利表').is_file())

    def test_docx_review_requires_nonempty_report_paragraph(self):
        source = self.docx('source.docx', ['原文'])
        for paragraphs in (['原文'], ['原文', '核查结果']):
            with self.subTest(paragraphs=paragraphs):
                output = self.docx('output.docx', paragraphs)
                with self.assertRaisesRegex(FileTaskError, '核查结果'):
                    self.run_task(source, output, '请核对文档')
        output = self.docx('output.docx', ['原文', '核查结果', '第二段日期无法核实。'])
        self.assertTrue(self.run_task(source, output, '请核对文档').is_file())
        output = self.docx('output.docx', ['原文', '核查结果：第二段日期无法核实。'])
        self.assertTrue(self.run_task(source, output, '请核对文档').is_file())

    def test_existing_report_is_preserved_and_a_new_report_is_required(self):
        source = self.xlsx('source.xlsx', report='上次核查结果')
        output = self.xlsx('output.xlsx', report='上次核查结果')
        with self.assertRaisesRegex(FileTaskError, '未新增'):
            self.run_task(source, output, '检查一下毛利表')
        book = Workbook()
        book.active['A1'] = '毛利表'
        book.create_sheet('核查结果')['A1'] = '上次核查结果'
        book.create_sheet('核查结果（2）')['A1'] = '本次核查范围与结论'
        book.save(output)
        self.assertTrue(self.run_task(source, output, '检查一下毛利表').is_file())

        book['核查结果（2）'].title = '核查结果(2)'
        book.save(output)
        self.assertTrue(self.run_task(source, output, '检查一下毛利表').is_file())

        book['核查结果'].title = '旧报告'
        book.save(output)
        with self.assertRaisesRegex(FileTaskError, '未保留原有工作表名称'):
            self.run_task(source, output, '检查一下毛利表')

    def test_ordinary_edit_does_not_require_review_report(self):
        source = self.xlsx('source.xlsx')
        self.assertFalse(is_review_request('把表头加粗，不需要检查'))
        self.assertTrue(self.run_task(source, source, '把表头加粗').is_file())

    def test_followup_uses_prior_result_and_preserves_original(self):
        source = self.xlsx('source.xlsx')
        original_bytes = source.read_bytes()
        previous = self.root / 'data/private/wxid_friend/files/previous/result.xlsx'
        previous.parent.mkdir(parents=True)
        book = Workbook()
        book.active['A1'] = '先前已经修改的数据'
        book.save(previous)
        def runner(folder, request, filename):
            shutil.copy2(folder / 'input' / filename, folder / 'output/result.xlsx')
        with patch('private_file_worker.locate_attachment', return_value=source), \
             patch('private_file_worker.check_text'):
            result = run_file_task(None, 'wxid_friend', {}, '帮忙修改', None, self.root,
                                   runner=runner, source_override=previous)
            self.assertEqual(load_workbook(result).active['A1'].value, '先前已经修改的数据')
            with self.assertRaisesRegex(FileTaskError, '无法安全继续'):
                run_file_task(None, 'wxid_friend', {}, '帮忙修改', None, self.root,
                              runner=runner, source_override=source)
        self.assertEqual(source.read_bytes(), original_bytes)


if __name__ == '__main__':
    unittest.main()
