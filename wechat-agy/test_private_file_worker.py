import shutil
import datetime
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook
from political_guard import PoliticalGuard
from private_file_worker import FileTaskError, office_text, run_file_task, codex_runner, locate_attachment


class Session:
    def __init__(self, answer='allow'):
        self.answer = answer
    def ask(self, payload):
        return {'decisions': [{'id': m['id'], 'decision': self.answer}
                              for m in payload['targets']], 'draft_decision': self.answer}
    def close(self):
        pass


class DB:
    def __init__(self, root):
        self.account_dir = str(root)
    def get_message_row(self, contact, local_id, shard):
        return {'local_type': 6 << 32 | 49, 'sender_username': contact,
                'create_time': datetime.datetime(2026, 10, 27).timestamp()}


class FileWorkerTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.folder = self.root / 'account' / 'msg' / 'file' / '2026-10'
        self.folder.mkdir(parents=True)
        self.db = DB(self.root / 'account')
        self.row = {'sender_username': 'wxid_test', 'shard': 'message_0.db', 'local_id': 1}
        self.guard = PoliticalGuard(model_session=Session())
        self.filename = '账目.xlsx'
    def tearDown(self):
        self.temp.cleanup()
    def xlsx(self, value='普通账目'):
        wb = Workbook(); wb.active['A1'] = value; wb.save(self.folder / '账目.xlsx')
    def copy_runner(self, folder, request, filename):
        shutil.copy2(folder / 'input' / filename, folder / 'output' / 'result.xlsx')
    def run_task(self, runner=None):
        filename = self.filename
        source = self.folder / filename
        size = source.stat().st_size if source.exists() else 10
        self.row['content'] = (f'<msg><appmsg><title>{filename}</title><type>6</type>'
                               f'<appattach><totallen>{size}</totallen></appattach></appmsg></msg>')
        self.row['create_time'] = datetime.datetime(2026, 10, 27).timestamp()
        return run_file_task(self.db, 'wxid_test', self.row, '整理表格', self.guard,
                             self.root, runner=runner or self.copy_runner)
    def test_xlsx_process_and_preserve_source(self):
        self.xlsx()
        result = self.run_task()
        self.assertTrue(result.exists())
        self.assertIn('普通账目', office_text(result))
        originals = list((self.root / 'data/private/wxid_test/files').rglob('original.xlsx'))
        self.assertEqual(len(originals), 1)
        self.assertEqual(originals[0].read_bytes(), (self.folder / '账目.xlsx').read_bytes())
    def test_missing_and_ambiguous(self):
        with self.assertRaisesRegex(FileTaskError, '尚未下载'): self.run_task()
        self.xlsx()
        other = self.folder.parent / '2026-09'; other.mkdir()
        shutil.copy2(self.folder / '账目.xlsx', other / '账目.xlsx')
        with self.assertRaisesRegex(FileTaskError, '同名'): self.run_task()
    def test_unknown_sender(self):
        self.xlsx(); self.row['sender_username'] = 'another'
        with self.assertRaisesRegex(FileTaskError, '身份'): self.run_task()

    def test_compound_file_type_rejects_incorrect_size(self):
        self.xlsx()
        self.row['create_time'] = datetime.datetime(2026, 10, 27).timestamp()
        self.row['content'] = ('<msg><appmsg><title>账目.xlsx</title><type>6</type>'
                               '<appattach><totallen>1</totallen></appattach></appmsg></msg>')
        with self.assertRaisesRegex(FileTaskError, '月份或大小不符'):
            locate_attachment(self.db, 'wxid_test', self.row)
    def test_political_input_and_uncertain_are_not_processed(self):
        self.xlsx('政治')
        with self.assertRaisesRegex(FileTaskError, '政治'): self.run_task()
        self.xlsx('普通内容'); self.guard = PoliticalGuard(model_session=Session('uncertain'))
        with self.assertRaisesRegex(FileTaskError, '政治'): self.run_task()
    def test_political_and_malformed_output_rejected(self):
        self.xlsx()
        def political(folder, request, filename):
            wb = Workbook(); wb.active['A1'] = '政治'; wb.save(folder / 'output/result.xlsx')
        with self.assertRaisesRegex(FileTaskError, '政治'): self.run_task(political)
        def malformed(folder, request, filename):
            (folder / 'output/result.xlsx').write_text('invalid')
        with self.assertRaisesRegex(FileTaskError, 'Office'): self.run_task(malformed)
    def test_docx(self):
        self.filename = '文档.docx'
        doc = Document(); doc.add_paragraph('日常事项'); doc.save(self.folder / '文档.docx')
        def copy(folder, request, filename):
            shutil.copy2(folder / 'input' / filename, folder / 'output/result.docx')
        result = self.run_task(copy)
        self.assertEqual(Document(result).paragraphs[0].text, '日常事项')

    @patch('private_file_worker.shutil.which', return_value='codex')
    @patch('private_file_worker.subprocess.run')
    def test_codex_zero_exit_without_output_is_failure(self, run, which):
        run.return_value.returncode = 0
        run.return_value.stdout = ''
        run.return_value.stderr = ''
        work = self.root / 'codex-test'
        (work / 'output').mkdir(parents=True)
        with self.assertRaisesRegex(FileTaskError, '未生成'):
            codex_runner(work, '处理', 'sample.xlsx')
        self.assertTrue((work / 'codex-diagnostic.json').exists())

    @patch('private_file_worker.shutil.which', return_value='codex')
    @patch('private_file_worker.subprocess.run')
    def test_codex_policy_rejection_is_not_success(self, run, which):
        run.return_value.returncode = 0
        run.return_value.stdout = ''
        run.return_value.stderr = 'command blocked by policy'
        work = self.root / 'codex-test'
        (work / 'output').mkdir(parents=True)
        with self.assertRaisesRegex(FileTaskError, '策略拒绝'):
            codex_runner(work, '处理', 'sample.xlsx')


if __name__ == '__main__':
    unittest.main()
